"""Synthetic *version* corpus and golden set v2 for the freshness/stale-docs eval.

The real notes have no superseded documents, so this generator builds document
*families*: the same project note in two or three versions whose planted facts
(budget, deadline, owner, technology, status) change between versions. Every
family follows one of six patterns, each exercising a different code path:

================  =====================================================  ==========
pattern           files                                                  linked by
================  =====================================================  ==========
same-folder       ``plano-orion.md`` -> ``plano-orion-v2.md``             name+content
archive           ``archive/plano-orion.md`` (old) -> ``plano-orion.md``  archive rule
date-stamped      ``2025-03-12-plano-orion.md`` -> ``2026-01-20-…``       name+content
frontmatter       ``plano-orion.md`` -> ``plano-orion-revisto.md``        ``supersedes:``
lookalike         two notes on the same project with different titles    nothing (both
                  (kick-off notes, then a status note)                   stay current)
in-place          one file rewritten in phase 2                          snapshots
================  =====================================================  ==========

Phase 1 (``--phase 1``, default) writes the initial files; phase 2
(``--phase 2``) rewrites the in-place families with their new facts, so the
runner can scan, apply the edits and scan again. ``manifest.json`` in the
output folder records every family, and ``--golden-out`` writes the golden set
whose questions are correct by construction: the newest version is the expected
document, older versions are ``stale_docs`` and their superseded values are
``stale_values``. Deterministic for a given ``--seed``; no LLM calls.

    uv run python scripts/eval/generate_version_corpus.py --out /tmp/mynotes-version-corpus \\
        --golden-out eval/golden_set_v2.yaml
"""

from __future__ import annotations

import argparse
import json
import os
import random
from dataclasses import asdict, dataclass, field
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any

import yaml

PATTERNS = ("same-folder", "archive", "date-stamped", "frontmatter", "lookalike", "in-place")
"""The family patterns, in generation order."""

LINKED_PATTERNS = frozenset({"same-folder", "archive", "date-stamped", "frontmatter"})
"""Patterns whose older versions become separate *superseded* documents."""

FACTS = ("budget", "deadline", "owner", "technology", "status")

_PROJECTS = [
    "Orion",
    "Vega",
    "Lince",
    "Tejo",
    "Alvor",
    "Sado",
    "Monchique",
    "Cintra",
    "Douro",
    "Minho",
    "Zezere",
    "Tamega",
    "Lima",
    "Cavado",
    "Mira",
    "Guadiana",
    "Arade",
    "Mondego",
    "Nabao",
    "Vouga",
    "Ave",
    "Leca",
    "Neiva",
    "Ancora",
    "Coa",
    "Sabor",
    "Tua",
    "Paiva",
    "Dao",
    "Alva",
    "Ceira",
    "Zambujal",
    "Odelouca",
    "Seda",
    "Caia",
    "Sorraia",
    "Almansor",
    "Raia",
    "Degebe",
    "Xarrama",
    "Ardila",
    "Cobres",
    "Terges",
    "Oeiras",
    "Lis",
    "Alviela",
    "Nabancia",
    "Sever",
]
_DOMAINS_PT = [
    "controlo de acessos do laboratorio",
    "monitorizacao de energia dos edificios",
    "gestao de inventario das oficinas",
    "portal de reservas de salas",
    "arquivo digital de projetos",
    "telemetria da frota",
    "assistente de pesquisa interna",
    "quadro de indicadores da producao",
    "catalogo de sensores",
    "plataforma de formacao interna",
    "gestao de pedidos de manutencao",
    "integracao com o ERP",
    "registo de visitantes",
    "painel de qualidade do ar",
]
_DOMAINS_EN = [
    "lab access control",
    "building energy monitoring",
    "workshop inventory management",
    "room booking portal",
    "digital project archive",
    "fleet telemetry",
    "internal search assistant",
    "production KPI dashboard",
    "sensor catalogue",
    "internal training platform",
    "maintenance request tracking",
    "ERP integration",
    "visitor registration",
    "air quality dashboard",
]
_PEOPLE = [
    "Ana Peixoto",
    "Rui Tavares",
    "Marta Carvalho",
    "Tiago Nunes",
    "Ines Barros",
    "Pedro Antunes",
    "Sofia Leitao",
    "Nuno Fonseca",
    "Carla Mendes",
    "Joao Ribeiro",
    "Beatriz Lopes",
    "Hugo Martins",
    "Rita Figueiredo",
    "Miguel Cordeiro",
    "Teresa Pinho",
    "Vasco Almeida",
    "Helena Guerreiro",
    "Bruno Sequeira",
    "Leonor Matos",
    "Diogo Pacheco",
]
_TECH = [
    "Rust",
    "Go",
    "Kotlin",
    "Elixir",
    "FastAPI",
    "Django",
    "Spring Boot",
    "PostgreSQL",
    "TimescaleDB",
    "Kafka",
    "RabbitMQ",
    "OpenSearch",
    "Qdrant",
    "React",
    "Svelte",
    "Flutter",
    "Grafana",
    "Airflow",
]
_STATUS_PT = ["em planeamento", "em desenvolvimento", "em piloto", "em producao", "suspenso"]
_STATUS_EN = ["in planning", "in development", "in pilot", "in production", "on hold"]
_MONTHS_PT = [
    "janeiro",
    "fevereiro",
    "marco",
    "abril",
    "maio",
    "junho",
    "julho",
    "agosto",
    "setembro",
    "outubro",
    "novembro",
    "dezembro",
]
_MONTHS_EN = [
    "January",
    "February",
    "March",
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
]

_WINDOWS = [
    (datetime(2025, 2, 1, tzinfo=UTC), datetime(2025, 6, 30, tzinfo=UTC)),
    (datetime(2025, 10, 1, tzinfo=UTC), datetime(2026, 2, 28, tzinfo=UTC)),
    (datetime(2026, 5, 1, tzinfo=UTC), datetime(2026, 8, 31, tzinfo=UTC)),
]
"""Document-date windows of version 1, 2 and 3."""


@dataclass
class Version:
    """One written version of a family."""

    number: int
    rel_path: str
    date: str  # YYYY-MM-DD
    facts: dict[str, str]
    title: str
    format: str
    phase: int = 1


@dataclass
class Family:
    """One document family: a project note in several versions."""

    id: str
    pattern: str
    language: str
    project: str
    slug: str
    domain: str
    folder: str
    changed_facts: list[str]
    constant_fact: str
    versions: list[Version] = field(default_factory=list)


# ---------------------------------------------------------------- facts and text


def _budget(rng: random.Random, language: str) -> str:
    amount = rng.randrange(12, 240) * 500
    text = f"{amount:,}".replace(",", " ")
    return f"{text} euros"


def _deadline(rng: random.Random, language: str, year: int) -> str:
    day = rng.randrange(1, 29)
    month = rng.randrange(12)
    if language == "pt":
        return f"{day} de {_MONTHS_PT[month]} de {year}"
    return f"{day} {_MONTHS_EN[month]} {year}"


def _fact_values(rng: random.Random, language: str, year: int) -> dict[str, str]:
    """A fresh, mutually distinct set of planted facts."""
    return {
        "budget": _budget(rng, language),
        "deadline": _deadline(rng, language, year),
        "owner": rng.choice(_PEOPLE),
        "technology": rng.choice(_TECH),
        "status": rng.choice(_STATUS_PT if language == "pt" else _STATUS_EN),
    }


def _next_facts(
    rng: random.Random, previous: dict[str, str], changed: list[str], language: str, year: int
) -> dict[str, str]:
    """The next version's facts: ``changed`` facts get a new, different value."""
    facts = dict(previous)
    for name in changed:
        for _ in range(50):
            candidate = _fact_values(rng, language, year)[name]
            if candidate != previous[name]:
                facts[name] = candidate
                break
    return facts


_FACT_SENTENCES = {
    "pt": {
        "budget": "O orcamento aprovado para o projeto e de {budget}.",
        "deadline": "O prazo de entrega acordado com a direcao e {deadline}.",
        "owner": "A pessoa responsavel pelo projeto e {owner}.",
        "technology": "A tecnologia escolhida para o backend e {technology}.",
        "status": "Estado atual do projeto: {status}.",
    },
    "en": {
        "budget": "The approved budget for the project is {budget}.",
        "deadline": "The delivery deadline agreed with management is {deadline}.",
        "owner": "The person responsible for the project is {owner}.",
        "technology": "The technology chosen for the backend is {technology}.",
        "status": "Current project status: {status}.",
    },
}

_BOILERPLATE = {
    "pt": {
        "Contexto": [
            "O Projeto {project} trata da {domain}. Esta nota resume o plano, a equipa e as "
            "decisoes tomadas ate a data indicada no cabecalho.",
            "A iniciativa partiu de um pedido da area operacional e foi avaliada pela comissao "
            "tecnica, que aprovou o arranque com as condicoes descritas abaixo.",
        ],
        "Objetivos": [
            "Reduzir o trabalho manual associado a {domain} e disponibilizar informacao fiavel "
            "em tempo util a quem dela precisa.",
            "Garantir que a solucao se integra com os sistemas existentes sem duplicar dados nem "
            "criar novos pontos de falha.",
        ],
        "Plano e orcamento": ["{budget_sentence}", "{deadline_sentence}", "{status_sentence}"],
        "Equipa e tecnologia": [
            "{owner_sentence}",
            "{technology_sentence}",
            "A equipa reune-se semanalmente e regista as decisoes nesta nota, que e a referencia "
            "para o estado do projeto.",
        ],
        "Riscos": [
            "O principal risco identificado e a dependencia de fornecedores externos para a "
            "componente de hardware.",
            "Existe tambem o risco de o calendario de ferias atrasar a fase de testes com os "
            "utilizadores finais.",
        ],
    },
    "en": {
        "Context": [
            "Project {project} covers {domain}. This note summarises the plan, the team and the "
            "decisions taken up to the date in the header.",
            "The initiative came from an operations request and was reviewed by the technical "
            "committee, which approved the start under the conditions below.",
        ],
        "Goals": [
            "Reduce the manual work around {domain} and make reliable information available in "
            "time to the people who need it.",
            "Make sure the solution integrates with the existing systems without duplicating "
            "data or adding new points of failure.",
        ],
        "Plan and budget": ["{budget_sentence}", "{deadline_sentence}", "{status_sentence}"],
        "Team and technology": [
            "{owner_sentence}",
            "{technology_sentence}",
            "The team meets weekly and records its decisions in this note, which is the "
            "reference for the project's state.",
        ],
        "Risks": [
            "The main risk identified is the dependency on external suppliers for the hardware "
            "component.",
            "There is also a risk that the holiday calendar delays the end-user testing phase.",
        ],
    },
}

_REVISION_NOTE = {
    "pt": "Nota de revisao: esta versao atualiza {items} face a versao anterior.",
    "en": "Revision note: this version updates {items} compared with the previous one.",
}
_FACT_LABEL = {
    "pt": {
        "budget": "o orcamento",
        "deadline": "o prazo",
        "owner": "a pessoa responsavel",
        "technology": "a tecnologia",
        "status": "o estado",
    },
    "en": {
        "budget": "the budget",
        "deadline": "the deadline",
        "owner": "the owner",
        "technology": "the technology",
        "status": "the status",
    },
}

_LOOKALIKE_FIRST = {
    "pt": (
        "Notas da reuniao de arranque - Projeto {project}",
        [
            (
                "Participantes",
                [
                    "Estiveram presentes a equipa tecnica, a area operacional e a direcao.",
                    "{owner_sentence}",
                ],
            ),
            (
                "Decisoes",
                [
                    "{budget_sentence}",
                    "{deadline_sentence}",
                    "{technology_sentence}",
                ],
            ),
            (
                "Proximos passos",
                [
                    "Preparar o ambiente de desenvolvimento e o plano de testes para a {domain}.",
                    "{status_sentence}",
                ],
            ),
        ],
    ),
    "en": (
        "Kick-off meeting notes - Project {project}",
        [
            (
                "Attendees",
                [
                    "The technical team, operations and management attended.",
                    "{owner_sentence}",
                ],
            ),
            (
                "Decisions",
                [
                    "{budget_sentence}",
                    "{deadline_sentence}",
                    "{technology_sentence}",
                ],
            ),
            (
                "Next steps",
                [
                    "Prepare the development environment and the test plan for {domain}.",
                    "{status_sentence}",
                ],
            ),
        ],
    ),
}
_LOOKALIKE_SECOND = {
    "pt": (
        "Ponto de situacao - Projeto {project}",
        [
            (
                "Resumo",
                [
                    "Ponto de situacao do Projeto {project} ({domain}) apos a revisao trimestral.",
                    "{status_sentence}",
                ],
            ),
            (
                "Alteracoes aprovadas",
                [
                    "{budget_sentence}",
                    "{deadline_sentence}",
                    "{technology_sentence}",
                ],
            ),
            ("Equipa", ["{owner_sentence}"]),
        ],
    ),
    "en": (
        "Status update - Project {project}",
        [
            (
                "Summary",
                [
                    "Status of Project {project} ({domain}) after the quarterly review.",
                    "{status_sentence}",
                ],
            ),
            (
                "Approved changes",
                [
                    "{budget_sentence}",
                    "{deadline_sentence}",
                    "{technology_sentence}",
                ],
            ),
            ("Team", ["{owner_sentence}"]),
        ],
    ),
}


def _fill(template: str, family: Family, facts: dict[str, str]) -> str:
    sentences = {
        f"{name}_sentence": _FACT_SENTENCES[family.language][name].format(**facts) for name in FACTS
    }
    return template.format(project=family.project, domain=family.domain, **sentences)


def _sections_for(family: Family, version: Version) -> list[tuple[str, list[str]]]:
    """The section list (title, paragraphs) of one version."""
    if family.pattern == "lookalike":
        _title, layout = (_LOOKALIKE_FIRST if version.number == 1 else _LOOKALIKE_SECOND)[
            family.language
        ]
        return [
            (title, [_fill(p, family, version.facts) for p in paragraphs])
            for title, paragraphs in layout
        ]
    sections = [
        (title, [_fill(p, family, version.facts) for p in paragraphs])
        for title, paragraphs in _BOILERPLATE[family.language].items()
    ]
    if version.number > 1:
        labels = [_FACT_LABEL[family.language][name] for name in family.changed_facts]
        last = " e " if family.language == "pt" else " and "
        items = labels[0] if len(labels) == 1 else ", ".join(labels[:-1]) + last + labels[-1]
        note = _REVISION_NOTE[family.language].format(items=items)
        sections.append(("Revisao" if family.language == "pt" else "Revision", [note]))
    return sections


# ---------------------------------------------------------------- file writers


def _set_mtime(path: Path, date: str) -> None:
    stamp = datetime.fromisoformat(date).replace(tzinfo=UTC, hour=10).timestamp()
    os.utime(path, (stamp, stamp))


def _write_markdown(
    path: Path, family: Family, version: Version, *, supersedes: str | None
) -> None:
    lines = [
        "---",
        f"title: {version.title}",
        f"date: {version.date}",
        "tags:",
        f"  - project:{family.slug}",
        "  - planning",
    ]
    if supersedes:
        lines.append(f"supersedes: {supersedes}")
    lines += ["---", "", f"# {version.title}", ""]
    for title, paragraphs in _sections_for(family, version):
        lines += [f"## {title}", "", *[p + "\n" for p in paragraphs]]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")
    _set_mtime(path, version.date)


def _write_docx(path: Path, family: Family, version: Version) -> None:
    import docx

    document = docx.Document()
    document.add_heading(version.title, level=0)
    for title, paragraphs in _sections_for(family, version):
        document.add_heading(title, level=1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    when = datetime.fromisoformat(version.date).replace(hour=10)
    document.core_properties.created = when
    document.core_properties.modified = when
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    _set_mtime(path, version.date)


def _write_pdf(path: Path, family: Family, version: Version) -> None:
    import fitz

    doc = fitz.open()
    stamp = datetime.fromisoformat(version.date).strftime("D:%Y%m%d100000")
    doc.set_metadata(
        {
            "title": version.title,
            "creationDate": stamp,
            "modDate": stamp,
            "producer": "mynotes-version-corpus",
            "creator": "mynotes-version-corpus",
        }
    )
    page = doc.new_page()
    y = 72.0
    page.insert_text((72, y), version.title, fontsize=16)
    y += 30
    for title, paragraphs in _sections_for(family, version):
        if y > 720:
            page = doc.new_page()
            y = 72.0
        page.insert_text((72, y), title, fontsize=12)
        y += 20
        for paragraph in paragraphs:
            rect = fitz.Rect(72, y, 523, y + 80)
            used = page.insert_textbox(rect, paragraph, fontsize=10)
            y += max(14.0, 80 - used) + 6
            if y > 720:
                page = doc.new_page()
                y = 72.0
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path), deflate=True)
    doc.close()
    _set_mtime(path, version.date)


def write_version(out: Path, family: Family, version: Version) -> None:
    """Write one version file in its format, with its supersedes link when the pattern needs it."""
    path = out / version.rel_path
    supersedes = None
    if family.pattern == "frontmatter" and version.number > 1:
        supersedes = Path(family.versions[version.number - 2].rel_path).name
    if version.format == "docx":
        _write_docx(path, family, version)
    elif version.format == "pdf":
        _write_pdf(path, family, version)
    else:
        _write_markdown(path, family, version, supersedes=supersedes)


# ---------------------------------------------------------------- planning


def _slugify(text: str) -> str:
    return text.lower().replace(" ", "-")


def _random_date(rng: random.Random, window: tuple[datetime, datetime]) -> str:
    start, end = window
    return (start + timedelta(days=rng.randrange((end - start).days + 1))).strftime("%Y-%m-%d")


def _version_paths(family: Family, count: int, fmt: str, dates: list[str]) -> list[str]:
    base = f"plano-{family.slug}" if family.language == "pt" else f"{family.slug}-plan"
    folder = family.folder
    if family.pattern == "same-folder":
        return [f"{folder}/{base}.{fmt}"] + [
            f"{folder}/{base}-v{n}.{fmt}" for n in range(2, count + 1)
        ]
    if family.pattern == "archive":
        # The newest keeps the plain name; every older one sits in the archive subfolder,
        # version 1 under its own name and the middle ones with a marker.
        paths = []
        for n in range(1, count):
            marker = "" if n == 1 else f"-v{n}"
            paths.append(f"{folder}/archive/{base}{marker}.{fmt}")
        return [*paths, f"{folder}/{base}.{fmt}"]
    if family.pattern == "date-stamped":
        return [f"{folder}/{date}-{base}.{fmt}" for date in dates]
    if family.pattern == "frontmatter":
        suffixes = [
            "",
            "-revisto" if family.language == "pt" else "-revised",
            "-atualizado" if family.language == "pt" else "-updated",
        ]
        return [f"{folder}/{base}{suffixes[n]}.{fmt}" for n in range(count)]
    if family.pattern == "lookalike":
        first = "notas-arranque" if family.language == "pt" else "kickoff-notes"
        second = "ponto-de-situacao" if family.language == "pt" else "status-update"
        return [f"{folder}/{first}-{family.slug}.{fmt}", f"{folder}/{second}-{family.slug}.{fmt}"]
    return [f"{folder}/{base}.{fmt}"]  # in-place: one path, rewritten


def plan_families(rng: random.Random, per_pattern: int) -> list[Family]:
    """Plan every family (paths, dates, facts) without writing anything."""
    projects = list(_PROJECTS)
    rng.shuffle(projects)
    families: list[Family] = []
    index = 0
    for pattern in PATTERNS:
        for k in range(per_pattern):
            project = projects[index]
            index += 1
            language = "pt" if k % 2 == 0 else "en"
            slug = _slugify(project)
            domain = rng.choice(_DOMAINS_PT if language == "pt" else _DOMAINS_EN)
            constant = rng.choice(["owner", "technology"])
            changeable = [f for f in FACTS if f != constant]
            changed = rng.sample(changeable, k=rng.choice([2, 3]))
            family = Family(
                id=f"f{len(families) + 1:02d}",
                pattern=pattern,
                language=language,
                project=project,
                slug=slug,
                domain=domain,
                folder=f"projects/{slug}",
                changed_facts=sorted(changed),
                constant_fact=constant,
            )
            # Versions: lookalike, in-place and archive have 2 (a third archived version
            # with a "-v2" marker would outrank the plain newest file by version number,
            # which is the detector's documented rule, not a detection failure); the other
            # linked patterns have 2 or 3.
            count = (
                2 if pattern in ("lookalike", "in-place", "archive") else (3 if k % 3 == 2 else 2)
            )
            fmt = "md"
            if pattern in ("same-folder", "archive") and k == per_pattern - 1:
                fmt = "docx" if pattern == "same-folder" else "pdf"
            dates = [_random_date(rng, _WINDOWS[n]) for n in range(count)]
            paths = _version_paths(family, count, fmt, dates)
            if pattern == "in-place":
                paths = paths * count
            facts = _fact_values(rng, language, 2026)
            for number in range(1, count + 1):
                if number > 1:
                    facts = _next_facts(rng, facts, family.changed_facts, language, 2026 + number)
                title = _sections_title(family, number)
                family.versions.append(
                    Version(
                        number=number,
                        rel_path=paths[number - 1],
                        date=dates[number - 1],
                        facts=dict(facts),
                        title=title,
                        format=fmt,
                        phase=2 if (pattern == "in-place" and number > 1) else 1,
                    )
                )
            families.append(family)
    return families


def _sections_title(family: Family, number: int) -> str:
    if family.pattern == "lookalike":
        layout = _LOOKALIKE_FIRST if number == 1 else _LOOKALIKE_SECOND
        return layout[family.language][0].format(project=family.project)
    if family.language == "pt":
        return f"Plano do Projeto {family.project}"
    return f"Project {family.project} plan"


# ---------------------------------------------------------------- distractors

_DISTRACTOR_TOPICS_PT = [
    ("Como configurar o acesso VPN", "ferramentas"),
    ("Checklist de onboarding", "pessoas"),
    ("Notas sobre backups semanais", "infraestrutura"),
    ("Guia de estilo para relatorios", "documentacao"),
    ("Resumo do workshop de seguranca", "formacao"),
    ("Procedimento de compras", "administracao"),
    ("Manutencao das impressoras", "infraestrutura"),
    ("Politica de ferias e ausencias", "pessoas"),
]
_DISTRACTOR_TOPICS_EN = [
    ("How to rotate API keys", "tools"),
    ("Monthly cost review", "administration"),
    ("Notes on the monitoring stack", "infrastructure"),
    ("Meeting room etiquette", "people"),
    ("Incident post-mortem template", "documentation"),
    ("Laptop refresh plan", "infrastructure"),
    ("Reading list on observability", "training"),
    ("Vendor evaluation criteria", "administration"),
]
_FILLER = {
    "pt": [
        "Esta nota foi escrita para registar o procedimento e evitar que o conhecimento fique "
        "apenas com uma pessoa.",
        "Sempre que o procedimento mudar, atualize esta nota e avise a equipa no canal habitual.",
        "Os passos abaixo foram testados em {tech} e devem funcionar nas versoes recentes.",
        "Em caso de duvida, fale com {person}, que acompanhou a ultima alteracao.",
        "O custo estimado desta atividade e de {budget} por ano, ja incluido no orcamento da area.",
    ],
    "en": [
        "This note records the procedure so the knowledge does not stay with one person only.",
        "Whenever the procedure changes, update this note and tell the team in the usual channel.",
        "The steps below were tested with {tech} and should work on recent versions.",
        "If in doubt, talk to {person}, who followed the last change.",
        "The estimated cost of this activity is {budget} per year, already in the area budget.",
    ],
}


def write_distractors(out: Path, rng: random.Random, count: int) -> list[str]:
    """Unrelated notes that share people, technologies and budget phrasing with the families."""
    paths: list[str] = []
    for i in range(count):
        language = "pt" if i % 2 == 0 else "en"
        topics = _DISTRACTOR_TOPICS_PT if language == "pt" else _DISTRACTOR_TOPICS_EN
        title, area = topics[i % len(topics)]
        fmt = "txt" if i % 9 == 8 else "md"
        rel = f"areas/{area}/{_slugify(title)}-{i + 1:02d}.{fmt}"
        path = out / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        date = _random_date(rng, rng.choice(_WINDOWS))
        body = [
            line.format(
                tech=rng.choice(_TECH), person=rng.choice(_PEOPLE), budget=_budget(rng, language)
            )
            for line in rng.sample(_FILLER[language], k=4)
        ]
        if fmt == "md":
            text = "\n".join(
                [
                    "---",
                    f"title: {title} {i + 1}",
                    f"date: {date}",
                    "---",
                    "",
                    f"# {title} {i + 1}",
                    "",
                    *[p + "\n" for p in body],
                ]
            )
        else:
            text = "\n".join([f"{title} {i + 1}", "", *body])
        path.write_text(text, encoding="utf-8")
        _set_mtime(path, date)
        paths.append(rel)
    return paths


# ---------------------------------------------------------------- golden set

_QUESTIONS = {
    "pt": {
        "budget": "Qual e o orcamento aprovado do Projeto {project}?",
        "deadline": "Qual e o prazo de entrega do Projeto {project}?",
        "owner": "Quem e a pessoa responsavel pelo Projeto {project}?",
        "technology": "Que tecnologia foi escolhida para o backend do Projeto {project}?",
        "status": "Qual e o estado atual do Projeto {project}?",
    },
    "en": {
        "budget": "What is the approved budget for Project {project}?",
        "deadline": "What is the delivery deadline for Project {project}?",
        "owner": "Who is the person responsible for Project {project}?",
        "technology": "Which technology was chosen for the backend of Project {project}?",
        "status": "What is the current status of Project {project}?",
    },
}
_ASOF_QUESTIONS = {
    "pt": {
        "budget": "Qual era o orcamento do Projeto {project}? as-of: {date}",
        "deadline": "Qual era o prazo de entrega do Projeto {project}? as-of: {date}",
        "technology": (
            "Que tecnologia estava escolhida para o backend do Projeto {project}? as-of: {date}"
        ),
        "status": "Qual era o estado do Projeto {project}? as-of: {date}",
        "owner": "Quem era a pessoa responsavel pelo Projeto {project}? as-of: {date}",
    },
    "en": {
        "budget": "What was the budget of Project {project}? as-of: {date}",
        "deadline": "What was the delivery deadline of Project {project}? as-of: {date}",
        "technology": "Which backend technology was chosen for Project {project}? as-of: {date}",
        "status": "What was the status of Project {project}? as-of: {date}",
        "owner": "Who was responsible for Project {project}? as-of: {date}",
    },
}


def _midpoint(a: str, b: str) -> str:
    start, end = datetime.fromisoformat(a), datetime.fromisoformat(b)
    return (start + (end - start) / 2).strftime("%Y-%m-%d")


def build_golden_set(
    families: list[Family], rng: random.Random, corpus_note: str
) -> dict[str, Any]:
    """The golden set v2: one current question per family, plus as-of and control questions."""
    questions: list[dict[str, Any]] = []
    counts: dict[str, int] = {}

    def add(
        category: str,
        family: Family,
        question: str,
        expected: list[str],
        answer: str,
        *,
        fact: str,
        stale_docs: list[str] | None = None,
        stale_values: list[str] | None = None,
    ) -> None:
        counts[category] = counts.get(category, 0) + 1
        row: dict[str, Any] = {
            "id": f"v{len(questions) + 1:03d}",
            "question": question,
            "language": family.language,
            "category": category,
            "tags": [f"pattern:{family.pattern}", f"fact:{fact}", f"family:{family.id}"],
            "filters": {"tags": [], "folders": []},
            "expected_docs": expected,
            "reference_answer": answer,
        }
        if stale_docs:
            row["stale_docs"] = stale_docs
        if stale_values:
            row["stale_values"] = stale_values
        questions.append(row)

    per_pattern_asof: dict[str, int] = {}
    per_pattern_control: dict[str, int] = {}
    for family in families:
        newest = family.versions[-1]
        older = family.versions[:-1]
        fact = rng.choice(family.changed_facts)
        old_values = sorted({v.facts[fact] for v in older} - {newest.facts[fact]})
        # Current-version question: the newest value must win.
        add(
            "current-version",
            family,
            _QUESTIONS[family.language][fact].format(project=family.project),
            [newest.rel_path],
            newest.facts[fact],
            fact=fact,
            stale_docs=[v.rel_path for v in older if v.rel_path != newest.rel_path] or None,
            stale_values=old_values or None,
        )
        # As-of question (two per linked pattern): the first version's value at a date before v2.
        if family.pattern in LINKED_PATTERNS and per_pattern_asof.get(family.pattern, 0) < 2:
            per_pattern_asof[family.pattern] = per_pattern_asof.get(family.pattern, 0) + 1
            first, second = family.versions[0], family.versions[1]
            as_of = _midpoint(first.date, second.date)
            add(
                "history",
                family,
                _ASOF_QUESTIONS[family.language][fact].format(project=family.project, date=as_of),
                [first.rel_path],
                first.facts[fact],
                fact=fact,
                stale_values=sorted(
                    {v.facts[fact] for v in family.versions[1:]} - {first.facts[fact]}
                )
                or None,
            )
        # Control question (one per pattern except lookalike): a fact that never changed.
        if family.pattern != "lookalike" and per_pattern_control.get(family.pattern, 0) < 1:
            per_pattern_control[family.pattern] = 1
            constant = family.constant_fact
            add(
                "control",
                family,
                _QUESTIONS[family.language][constant].format(project=family.project),
                [newest.rel_path],
                newest.facts[constant],
                fact=constant,
            )
    return {
        "meta": {
            "created": datetime.now(UTC).strftime("%Y-%m-%d"),
            "corpus_note": corpus_note,
            "review": "GENERATED (not human-reviewed): questions are correct by construction from "
            "the synthetic corpus manifest; review the wording, not the answers.",
            "counts": counts,
        },
        "questions": questions,
    }


# ---------------------------------------------------------------- main


def generate(
    out: Path, *, seed: int, per_pattern: int, distractors: int, phase: str, golden_out: Path | None
) -> dict[str, Any]:
    """Generate the corpus (phase 1 and/or 2) and optionally the golden set."""
    manifest_path = out / "manifest.json"
    if phase in ("1", "all"):
        if out.exists():
            for child in sorted(out.rglob("*"), reverse=True):
                if child.is_file() or child.is_symlink():
                    child.unlink()
                elif child.is_dir():
                    child.rmdir()
        out.mkdir(parents=True, exist_ok=True)
        rng = random.Random(seed)
        families = plan_families(rng, per_pattern)
        for family in families:
            for version in family.versions:
                if version.phase == 1:
                    write_version(out, family, version)
        distractor_paths = write_distractors(out, rng, distractors)
        manifest = {
            "seed": seed,
            "families": [asdict(f) for f in families],
            "distractors": distractor_paths,
        }
        manifest_path.write_text(
            json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        if golden_out is not None:
            note = (
                f"synthetic version corpus: {len(families)} families "
                f"({per_pattern} per pattern), {distractors} distractors, seed {seed}"
            )
            golden = build_golden_set(families, random.Random(seed + 1), note)
            header = (
                "# Golden set v2 - synthetic version corpus\n"
                "# (scripts/eval/generate_version_corpus.py).\n"
                "# GENERATED, not human-reviewed. Answers are correct by construction; review the\n"
                "# wording. Regenerate with the same --seed to reproduce. Never edit between two\n"
                "# runs you want to compare.\n"
            )
            golden_out.write_text(
                header + yaml.safe_dump(golden, sort_keys=False, allow_unicode=True, width=100),
                encoding="utf-8",
            )
    if phase in ("2", "all"):
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        families = [
            Family(**{**f, "versions": [Version(**v) for v in f["versions"]]})
            for f in manifest["families"]
        ]
        edited = 0
        for family in families:
            for version in family.versions:
                if version.phase == 2:
                    write_version(out, family, version)
                    edited += 1
        manifest["phase2_edits"] = edited
        manifest_path.write_text(
            json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8"
        )
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    files = [p for p in out.rglob("*") if p.is_file() and p.name != "manifest.json"]
    return {
        "out": str(out),
        "families": len(manifest["families"]),
        "files": len(files),
        "by_format": {
            ext: sum(1 for p in files if p.suffix == ext)
            for ext in sorted({p.suffix for p in files})
        },
        "phase2_edits": manifest.get("phase2_edits"),
        "golden_out": str(golden_out) if golden_out else None,
    }


def main() -> None:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("--out", type=Path, default=Path("/tmp/mynotes-version-corpus"))
    parser.add_argument(
        "--golden-out",
        type=Path,
        default=None,
        help="write the golden set v2 YAML here (phase 1 only)",
    )
    parser.add_argument("--seed", type=int, default=2)
    parser.add_argument("--families-per-pattern", type=int, default=7)
    parser.add_argument("--distractors", type=int, default=60)
    parser.add_argument(
        "--phase",
        choices=["1", "2", "all"],
        default="1",
        help="1 = initial files (+ golden set), 2 = in-place edits, all = both",
    )
    args = parser.parse_args()
    summary = generate(
        args.out,
        seed=args.seed,
        per_pattern=args.families_per_pattern,
        distractors=args.distractors,
        phase=args.phase,
        golden_out=args.golden_out,
    )
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
