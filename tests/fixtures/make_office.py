"""Fixture-generation helpers for Office/HTML/EPUB/OCR tests.

No binary fixtures are committed: every helper writes its file into the
git-ignored ``tests/fixtures/docs/`` directory (or a caller-provided path).
The DOCX/PPTX/XLSX fixtures use python-docx/python-pptx/openpyxl (dev
dependencies); the EPUB is a hand-rolled minimal EPUB 2 zip so the EPUB
extraction path is always testable without ebooklib.
"""

from __future__ import annotations

import zipfile
from datetime import UTC, datetime
from pathlib import Path
from xml.sax.saxutils import escape

DOCS = Path(__file__).resolve().parent / "docs"

DOCX_TITLE = "Proposta Alfa"
PPTX_SLIDES = (
    ("Visão", "Este projeto cria um assistente local.", None),
    ("Plano", "Primeiro construímos o índice.", "Lembrar de falar do orçamento."),
    ("Riscos", "Dependência de um único endpoint.", None),
)
XLSX_SHEETS = (
    ("Custos", (("Item", "Valor"), ("Filtro", 10))),
    ("Prazos", (("Fase", "Mês"), ("Setup", 1))),
)
HTML_NAV_SOURCE = (
    "<html><head><script>x()</script></head><body>"
    '<nav><a href="/inicio">Início</a> <a href="/sobre">Sobre</a></nav>'
    "<h1>Página</h1><p>Olá mundo</p></body></html>"
)
EPUB_TITLE = "Guia de Bolso"
EPUB_CHAPTERS = (
    ("Capítulo Um", "Este é o primeiro capítulo do guia de bolso."),
    ("Capítulo Dois", "Segundo capítulo com conteúdo."),
)
OCR_TEXT = "Filtro do exaustor deve ser limpo todo mes"


def make_docx(path: Path | None = None) -> Path:
    """Write the sample DOCX: H1 + H2 headings, a paragraph and a 2x3 table."""
    from docx import Document

    target = path or (DOCS / "sample.docx")
    document = Document()
    document.add_heading(DOCX_TITLE, level=1)
    document.add_heading("Orçamento", level=2)
    document.add_paragraph("Este documento descreve a proposta alfa em detalhe.")
    table = document.add_table(rows=2, cols=3)
    rows = (("Item", "Qtd", "Preço"), ("Filtro", "2", "10,00"))
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            table.rows[row_index].cells[col_index].text = value
    document.core_properties.title = DOCX_TITLE
    document.core_properties.modified = datetime(2026, 3, 1, 12, 0, tzinfo=UTC)
    document.save(target)
    return target


def make_pptx(path: Path | None = None) -> Path:
    """Write the sample PPTX: three titled slides, speaker notes on slide 2."""
    from pptx import Presentation

    target = path or (DOCS / "sample.pptx")
    prs = Presentation()
    for title, body, notes in PPTX_SLIDES:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    prs.core_properties.title = "Apresentação do projeto"
    prs.save(target)
    return target


def make_xlsx(path: Path | None = None) -> Path:
    """Write the sample XLSX with the two sheets Custos and Prazos."""
    from openpyxl import Workbook

    target = path or (DOCS / "sample.xlsx")
    wb = Workbook()
    sheet = wb.active
    assert sheet is not None
    first_name, first_rows = XLSX_SHEETS[0]
    sheet.title = first_name
    for row in first_rows:
        sheet.append(row)
    for sheet_name, rows in XLSX_SHEETS[1:]:
        created = wb.create_sheet(sheet_name)
        for row in rows:
            created.append(row)
    wb.save(target)
    return target


def make_html(path: Path | None = None, source: str = HTML_NAV_SOURCE) -> Path:
    """Write the sample HTML: a script, navigation links and one heading."""
    target = path or (DOCS / "sample.html")
    target.write_text(source, encoding="utf-8")
    return target


def make_epub(path: Path | None = None) -> Path:
    """Write a minimal EPUB 2 (mimetype + container + OPF + two XHTML chapters)."""
    target = path or (DOCS / "sample.epub")
    container = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
        '<rootfiles><rootfile full-path="OEBPS/content.opf" '
        'media-type="application/oebps-package+xml"/></rootfiles></container>'
    )
    manifest = "".join(
        f'<item id="c{index}" href="ch{index}.xhtml" media-type="application/xhtml+xml"/>'
        for index in range(1, len(EPUB_CHAPTERS) + 1)
    )
    spine = "".join(f'<itemref idref="c{index}"/>' for index in range(1, len(EPUB_CHAPTERS) + 1))
    opf = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<package xmlns="http://www.idpf.org/2007/opf" version="2.0" unique-identifier="id">'
        '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
        f"<dc:title>{escape(EPUB_TITLE)}</dc:title>"
        "<dc:language>pt</dc:language>"
        '<dc:identifier id="id">urn:uuid:mynotes-fixture</dc:identifier>'
        "</metadata>"
        f"<manifest>{manifest}</manifest>"
        f"<spine>{spine}</spine>"
        "</package>"
    )
    with zipfile.ZipFile(target, "w") as archive:
        archive.writestr("mimetype", "application/epub+zip")
        archive.writestr("META-INF/container.xml", container)
        archive.writestr("OEBPS/content.opf", opf)
        for index, (heading, paragraph) in enumerate(EPUB_CHAPTERS, start=1):
            archive.writestr(
                f"OEBPS/ch{index}.xhtml",
                '<?xml version="1.0" encoding="UTF-8"?>'
                '<html xmlns="http://www.w3.org/1999/xhtml"><head><title>c'
                f"{index}</title></head><body><h1>{escape(heading)}</h1>"
                f"<p>{escape(paragraph)}</p></body></html>",
            )
    return target


def make_image_only_pdf(path: Path, text: str = OCR_TEXT) -> Path:
    """Write a PDF whose only content is a rasterized image of ``text``."""
    import pymupdf

    staged = pymupdf.open()
    try:
        page = staged.new_page()
        page.insert_text((72, 100), text, fontsize=18, fontname="helv")
        pixmap = page.get_pixmap(dpi=200)
        png = pixmap.tobytes("png")
        width, height = pixmap.width, pixmap.height
    finally:
        staged.close()
    document = pymupdf.open()
    try:
        page = document.new_page(width=width, height=height)
        page.insert_image(page.rect, stream=png)
        document.save(path)
    finally:
        document.close()
    return path
