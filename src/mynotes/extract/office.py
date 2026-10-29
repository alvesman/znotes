"""Office, HTML and EPUB extractor built on markitdown (PLAN.md §7.2).

markitdown turns DOCX, PPTX, HTML and EPUB into Markdown with headings
preserved (XLSX sheets as ``## SheetName`` + table, PPTX slides marked with
``<!-- Slide number: N -->`` and speaker notes under ``### Notes:``). This
module feeds that Markdown into the shared ``markdown_to_blocks`` IR builder,
splitting slides into per-slide block groups, and resolves title and document
date from the OOXML ``docProps/core.xml`` read straight out of the zip (no
extra runtime dependency). Legacy ``.doc``/``.ppt`` are out of scope (PLAN.md
§1): the scanner records them as unsupported, no extractor is registered.
"""

from __future__ import annotations

import logging
import re
import xml.etree.ElementTree as ET
import zipfile
from datetime import UTC, datetime
from pathlib import Path

from markitdown import MarkItDown

from mynotes.extract.dates import resolve_doc_date
from mynotes.extract.ir import IRBlock, IRDocument, Link
from mynotes.extract.markdown import markdown_to_blocks
from mynotes.extract.registry import register
from mynotes.extract.titles import resolve_title
from mynotes.text.language import detect_language

logger = logging.getLogger(__name__)

_MARKITDOWN: MarkItDown | None = None

FORMATS: dict[str, str] = {
    "docx": "docx",
    "pptx": "pptx",
    "xlsx": "xlsx",
    "html": "html",
    "htm": "html",
    "epub": "epub",
}
"""File extension (without dot) to scanner format name handled here."""

_SLIDE_MARKER_RE = re.compile(r"^<!--\s*Slide number:\s*(\d+)\s*-->\s*$")
"""markitdown's per-slide marker line in PPTX output."""

_LINK_SPAN_RE = re.compile(r"\[[^\]]*\]\([^)]*\)")
_WORD_CHAR_RE = re.compile(r"[^\W_]", re.UNICODE)
"""Word characters without the underscore: leftover letters after link removal."""

_FILE_LIKE_TITLE_RE = re.compile(r"\.[A-Za-z0-9]{1,5}$")
"""Titles ending in a short extension (``.docx``, ``.indd``) look like file names."""

_PLACEHOLDER_TITLES = frozenset({"document", "documento", "presentation", "apresentação"})
"""Core-properties titles that carry no information (case-insensitive)."""

_CORE_XML = "docProps/core.xml"
_CP_NS = "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
_DC_NS = "http://purl.org/dc/elements/1.1/"
_DCTERMS_NS = "http://purl.org/dc/terms/"


def get_markitdown() -> MarkItDown:
    """Return the shared :class:`MarkItDown` instance, created on first use."""
    global _MARKITDOWN
    if _MARKITDOWN is None:
        _MARKITDOWN = MarkItDown()
    return _MARKITDOWN


class OfficeExtractor:
    """Extractor for DOCX, PPTX, XLSX, HTML and EPUB via markitdown.

    XLSX is read with openpyxl directly, in the same Markdown shape markitdown
    gives (a ``##`` heading per sheet, then a table): markitdown's own XLSX
    path goes through pandas, 72 MB the image does not otherwise need.
    """

    name = "markitdown"
    version = "1"
    formats = ("docx", "pptx", "xlsx", "html", "epub")

    def extract(self, path: Path) -> IRDocument:
        """Extract an Office/HTML/EPUB file into the common IR.

        Args:
            path: The file to read.

        Returns:
            The IR document with per-slide blocks for PPTX (``page`` = slide
            number, speaker notes as blocks of their slide), sheet headings
            and tables for XLSX, and plain Markdown blocks otherwise; title,
            document date (Office core properties when present) and language
            are resolved.
        """
        fmt = FORMATS.get(path.suffix.lower().lstrip("."), path.suffix.lower().lstrip("."))
        if fmt == "xlsx":
            markdown = xlsx_to_markdown(path)
        else:
            markdown = get_markitdown().convert(str(path)).markdown
        logger.debug("extraction produced %d characters for %s", len(markdown), path.name)

        if fmt == "pptx":
            blocks, links, first_slide_title = _pptx_blocks(markdown)
        elif fmt == "html":
            html_blocks, links = markdown_to_blocks(markdown)
            blocks = _strip_leading_nav_links(html_blocks)
            first_slide_title = None
        else:
            blocks, links = markdown_to_blocks(markdown)
            first_slide_title = None

        metadata_title: str | None = None
        metadata_date: datetime | None = None
        if fmt in ("docx", "pptx"):
            title, modified, created = read_core_properties(path)
            metadata_title = _clean_metadata_title(title)
            metadata_date = modified or created

        document = IRDocument(
            blocks=blocks,
            frontmatter={},
            links=links,
            hashtags=[],
            extractor=self.name,
            extractor_version=self.version,
        )
        document.language = detect_language(document.plain_text()[:5000])
        document.title = resolve_title(
            frontmatter={},
            blocks=blocks,
            filename=path.name,
            first_slide_title=first_slide_title,
            metadata_title=metadata_title,
        )
        document.doc_date = resolve_doc_date(
            frontmatter={},
            filename=path.name,
            metadata_date=metadata_date,
            mtime_ns=path.stat().st_mtime_ns,
        )
        return document


def xlsx_to_markdown(path: Path) -> str:
    """Render a workbook as Markdown: a ``##`` heading per sheet, then its rows as a table.

    The first non-empty row is the header. Empty rows are skipped; formulas
    give their cached value; ``|`` and line breaks inside a cell are escaped
    or collapsed so the table stays one row per line. A sheet without content
    gives only its heading.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        parts: list[str] = []
        for sheet in workbook.worksheets:
            rows = [
                [_cell_text(value) for value in row] for row in sheet.iter_rows(values_only=True)
            ]
            rows = [row for row in rows if any(row)]
            lines = [f"## {sheet.title}"]
            if rows:
                width = max(len(row) for row in rows)
                padded = [row + [""] * (width - len(row)) for row in rows]
                lines.append("| " + " | ".join(padded[0]) + " |")
                lines.append("| " + " | ".join(["---"] * width) + " |")
                lines.extend("| " + " | ".join(row) + " |" for row in padded[1:])
            parts.append("\n".join(lines))
        return "\n\n".join(parts) + "\n"
    finally:
        workbook.close()


def _cell_text(value: object) -> str:
    """One cell as table text: None empty, whole floats without ``.0``, no bars or newlines."""
    if value is None:
        return ""
    text = str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)
    return " ".join(text.split()).replace("|", "\\|")


def _pptx_blocks(markdown: str) -> tuple[list[IRBlock], list[Link], str | None]:
    """Split markitdown's PPTX Markdown on slide markers.

    Args:
        markdown: The markitdown output.

    Returns:
        ``(blocks, links, first_slide_title)`` where every block carries its
        slide number in ``page`` (content before the first marker keeps
        ``page=None``) and the first heading of slide 1 is the title fallback.
    """
    blocks: list[IRBlock] = []
    links: list[Link] = []
    first_slide_title: str | None = None
    for page, part in _split_slides(markdown):
        slide_blocks, slide_links = markdown_to_blocks(part, page=page)
        blocks.extend(slide_blocks)
        links.extend(slide_links)
        if page == 1:
            first_slide_title = next(
                (block.text for block in slide_blocks if block.kind == "heading"), None
            )
    return blocks, links, first_slide_title


def _split_slides(markdown: str) -> list[tuple[int | None, str]]:
    """Split PPTX Markdown into ``(slide_number, part)`` pairs.

    Content before the first slide marker (defensive; markitdown always starts
    with slide 1) is returned with ``slide_number=None``.
    """
    parts: list[tuple[int | None, str]] = []
    preamble: list[str] = []
    slide: tuple[int, list[str]] | None = None
    for line in markdown.splitlines():
        match = _SLIDE_MARKER_RE.match(line.strip())
        if match:
            if slide is not None:
                parts.append((slide[0], "\n".join(slide[1])))
            elif any(part.strip() for part in preamble):
                parts.append((None, "\n".join(preamble)))
            slide = (int(match.group(1)), [])
        elif slide is not None:
            slide[1].append(line)
        else:
            preamble.append(line)
    if slide is not None:
        parts.append((slide[0], "\n".join(slide[1])))
    return parts


def _strip_leading_nav_links(blocks: list[IRBlock]) -> list[IRBlock]:
    """Drop blocks before the first heading whose text is only link spans.

    HTML boilerplate navigation renders in markitdown as paragraphs or bullet
    lists made purely of ``[text](url)`` spans; real content is kept.

    Args:
        blocks: Blocks parsed from the HTML Markdown.

    Returns:
        The blocks with pure-link lead-ins removed.
    """
    kept: list[IRBlock] = []
    seen_heading = False
    for block in blocks:
        if block.kind == "heading":
            seen_heading = True
        if not seen_heading and block.kind in ("paragraph", "list") and _is_pure_links(block.text):
            logger.debug("Dropping navigation link block: %r", block.text[:80])
            continue
        kept.append(block)
    return kept


def _is_pure_links(text: str) -> bool:
    """Return True when ``text`` holds at least one link span and no other words."""
    if not _LINK_SPAN_RE.search(text):
        return False
    remainder = _LINK_SPAN_RE.sub("", text)
    return _WORD_CHAR_RE.search(remainder) is None


def _clean_metadata_title(value: str | None) -> str | None:
    """Return the core-properties title unless it is empty or a placeholder.

    Args:
        value: Raw ``dc:title`` from ``docProps/core.xml``.

    Returns:
        The stripped title, or None when it is empty, ends with an extension
        (file-name-like, e.g. ``823245pt.indd``) or is one of the known
        placeholder words ("Document", "Documento", "Presentation",
        "Apresentação"), compared case-insensitively.
    """
    if value is None:
        return None
    cleaned = value.strip()
    if not cleaned:
        return None
    if _FILE_LIKE_TITLE_RE.search(cleaned):
        return None
    if cleaned.lower() in _PLACEHOLDER_TITLES:
        return None
    return cleaned


def read_core_properties(path: Path) -> tuple[str | None, datetime | None, datetime | None]:
    """Read OOXML core properties straight from the zip (no python-docx/pptx).

    Args:
        path: A DOCX or PPTX file.

    Returns:
        ``(title, modified, created)`` from ``docProps/core.xml``; every value
        is None when the part is missing or unreadable. Datetimes are
        timezone-aware UTC (naive values are assumed UTC).
    """
    try:
        with zipfile.ZipFile(path) as archive:
            raw = archive.read(_CORE_XML)
    except (KeyError, zipfile.BadZipFile, OSError):
        return None, None, None
    try:
        root = ET.fromstring(raw)
    except ET.ParseError:
        logger.debug("Malformed %s in %s", _CORE_XML, path.name)
        return None, None, None
    title = _find_text(root, f"{{{_DC_NS}}}title")
    return (
        title,
        _find_date(root, f"{{{_DCTERMS_NS}}}modified"),
        _find_date(root, f"{{{_DCTERMS_NS}}}created"),
    )


def _find_text(root: ET.Element, tag: str) -> str | None:
    """Return the stripped text of the first matching element, or None."""
    element = root.find(tag)
    if element is None or element.text is None:
        return None
    cleaned = element.text.strip()
    return cleaned or None


def _find_date(root: ET.Element, tag: str) -> datetime | None:
    """Parse a ``dcterms`` W3CDTF element into an aware UTC datetime."""
    element = root.find(tag)
    if element is None or not element.text:
        return None
    try:
        parsed = datetime.fromisoformat(element.text.strip().replace("Z", "+00:00"))
    except ValueError:
        return None
    if parsed.tzinfo is None:
        return parsed.replace(tzinfo=UTC)
    return parsed.astimezone(UTC)


register(OfficeExtractor())
